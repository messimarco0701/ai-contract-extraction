"""
Runs the whole flow on a folder of documents.

  1. text.py turns each PDF or Excel file into plain text
  2. an extractor (Claude or the rule-based baseline) fills a ContractFields record
  3. validation.py normalizes the record and lists every reason not to trust it
  4. records with no reasons are "ok", everything else is "needs_review"
  5. the result goes to an Excel file with the full table and a review queue,
     plus a CSV of the same table for loading into another system
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import yaml
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .extractors import Extractor
from .schema import FIELDS
from .text import find_documents, read_text
from .validation import normalize, validate

ROOT = Path(__file__).resolve().parents[1]


def load_config(path: Path = ROOT / "config.yaml") -> dict:
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def validation_rules(config: dict) -> dict:
    return {**config["validation"], "own_vat_number": config["own_company"]["vat_number"]}


def process_document(path: Path, extractor: Extractor, rules: dict) -> dict:
    text = read_text(path)
    result = extractor.extract(text)
    record = normalize(result.fields)
    issues = validate(record, rules, source_text=text)
    if result.error:
        issues.insert(0, result.error)
    return {
        "doc_id": path.stem,
        "file": path.name,
        "extractor": extractor.name,
        "status": "needs_review" if issues else "ok",
        "issues": "; ".join(issues),
        **record.model_dump(),
    }


def run(input_dir: Path, extractor: Extractor, config: dict) -> pd.DataFrame:
    rules = validation_rules(config)
    rows = [process_document(p, extractor, rules) for p in find_documents(input_dir)]
    df = pd.DataFrame(rows, columns=["doc_id", "file", "extractor", "status", "issues", *FIELDS])
    return df


def write_outputs(df: pd.DataFrame, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    df.to_csv(out_dir / "extracted.csv", index=False, encoding="utf-8")
    review = df[df["status"] == "needs_review"][["doc_id", "file", "issues", *FIELDS]]
    path = out_dir / "extracted.xlsx"
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    review_fill = PatternFill("solid", fgColor="FCE4D6")
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, sheet in {"Extracted": df, "Review queue": review}.items():
            sheet.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
            for i, col in enumerate(sheet.columns, start=1):
                width = max([len(str(col))] + [len(str(v)) for v in sheet[col]])
                ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
            if name == "Extracted":
                status_col = list(sheet.columns).index("status") + 1
                for row in range(2, len(sheet) + 2):
                    if ws.cell(row=row, column=status_col).value == "needs_review":
                        ws.cell(row=row, column=status_col).fill = review_fill
    return path
