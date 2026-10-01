"""
Turns a document into plain text, whatever format it arrived in.

Both extractors work on text, not on files. That keeps them interchangeable and
makes the pipeline easy to extend: a new input format only needs a new reader
here. Scanned images would need OCR at this step; the samples in this repo all
have a text layer, so OCR is out of scope.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

SUPPORTED = {".pdf", ".xlsx"}


def _cell(value) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def pdf_text(path: Path) -> str:
    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def xlsx_text(path: Path) -> str:
    """
    One line per row, cells separated by ' | '. Layout is left for the extractor to make sense of.

    Empty cells inside a row are kept as empty slots. The first version dropped
    them, and on export-style sheets (headers in one row, values in the next)
    every value after a blank cell slid under the wrong header.
    """
    wb = load_workbook(path, data_only=True, read_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else _cell(v).strip() for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if cells:
                lines.append(" | ".join(cells))
    wb.close()
    return "\n".join(lines)


def read_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return pdf_text(path)
    if suffix == ".xlsx":
        return xlsx_text(path)
    raise ValueError(f"Unsupported file type: {path.name}")


def find_documents(folder: Path) -> list[Path]:
    return sorted(p for p in folder.iterdir() if p.suffix.lower() in SUPPORTED and not p.name.startswith("~$"))
